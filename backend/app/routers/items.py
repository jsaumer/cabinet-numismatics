import csv
import io
import json
import uuid
from datetime import datetime, timezone
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.auth import events
from app.auth.permissions import permission, require
from app.db import get_db
from app.models import CatalogRef, Grade, Item, ItemEvent, ItemSet, PriceEstimate, Tag
from app.schemas import (
    YEAR_REQUIRED,
    BulkResult,
    BulkUpdate,
    EventOut,
    ItemCreate,
    ItemDetail,
    ItemList,
    ItemListEntry,
    ItemOut,
    ItemUpdate,
    RunCreate,
    RunResult,
    SimilarItem,
)
from app.services import (
    app_settings,
    calendars,
    checklists,
    duplicates,
    importing,
    numista,
    pricing,
    serials,
    trash,
)
from app.services.currency import Converter

router = APIRouter(prefix="/api/items", tags=["items"])

SORTABLE = {
    "created_at": Item.created_at,
    "year": Item.year,
    "country": Item.country,
    "denomination": Item.denomination,
    "acquisition_date": Item.acquisition_date,
    "acquisition_price": Item.acquisition_price,
    "priority": Item.priority,
    "target_price": Item.target_price,
    "pcgs_pop_higher": Item.pcgs_pop_higher,
}
# Mostly empty columns: rows without a value sort last in either direction
# (an undated piece has no year).
NULLS_LAST = {"priority", "target_price", "year", "pcgs_pop_higher", "value"}

ITEM_LOAD = (
    selectinload(Item.photos),
    selectinload(Item.estimates),
    selectinload(Item.tags),
    selectinload(Item.catalog_refs),
    selectinload(Item.comparables),
    selectinload(Item.documents),
)

CSV_COLUMNS = [
    "id",
    "type",
    "status",
    "country",
    "denomination",
    "year",
    "year_nd",
    "mint_mark",
    "series",
    "variety",
    "strike",
    "composition",
    "weight_g",
    "fineness",
    "diameter_mm",
    "thickness_mm",
    "width_mm",
    "height_mm",
    "edge",
    "shape",
    "mintage",
    "die_axis",
    "struck_calendar",
    "struck_year",
    "struck_era",
    "grade_scale",
    "grade",
    "grade_plus",
    "grade_star",
    "designations",
    "grade_details",
    "cac_sticker",
    "cert_service",
    "cert_number",
    "pcgs_population",
    "pcgs_pop_higher",
    "serial_number",
    "prefix_block",
    "signatures",
    "issuer",
    "replacement_note",
    "charter_number",
    "bank_city",
    "bank_state",
    "plate_position",
    "printer",
    "watermark",
    "demonetized_on",
    "quantity",
    "acquisition_date",
    "acquisition_price",
    "acquisition_fees",
    "spot_at_purchase",
    "currency",
    "acquired_from",
    "storage_location",
    "sold_date",
    "sold_price",
    "sold_fees",
    "sold_to",
    "target_price",
    "priority",
    "notes",
    "tags",
    "catalog_refs",
    "set",
    "custom_fields",
    "latest_value",
    "latest_value_currency",
    "created_at",
]

# Import accepts the export format; these columns are derived and ignored on the way in.
IMPORT_IGNORED = {"id", "latest_value", "latest_value_currency", "created_at"}

# Grade prefixes a label may use for a proof or specimen strike on the Sheldon scale.
STRIKE_PREFIXES = {"PR": "proof", "PF": "proof", "SP": "specimen"}


def get_item_or_404(
    db: Session, item_id: uuid.UUID, *, load_related: bool = False, include_deleted: bool = False
) -> Item:
    """The item, or 404. Items in the trash count as missing (so they can't
    be edited) unless `include_deleted` (viewing, restoring, purging)."""
    stmt = select(Item).where(Item.id == item_id)
    if include_deleted:
        stmt = stmt.execution_options(include_deleted=True)
    if load_related:
        stmt = stmt.options(*ITEM_LOAD)
    item = db.execute(stmt).scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Item not found")
    return item


def resolve_tags(db: Session, names: list[str]) -> list[Tag]:
    """Get-or-create tags by name (case-preserving, whitespace-trimmed)."""
    tags = []
    for name in dict.fromkeys(n.strip() for n in names if n.strip()):
        tag = db.execute(select(Tag).where(Tag.name == name)).scalar_one_or_none()
        if tag is None:
            tag = Tag(name=name)
            db.add(tag)
            db.flush()
        tags.append(tag)
    return tags


def resolve_catalog_refs(db: Session, refs) -> list[CatalogRef]:
    resolved = []
    seen = set()
    for ref in refs:
        key = (ref.catalog.strip().lower(), ref.ref_code.strip())
        if not key[0] or not key[1] or key in seen:
            continue
        seen.add(key)
        row = db.execute(
            select(CatalogRef).where(CatalogRef.catalog == key[0], CatalogRef.ref_code == key[1])
        ).scalar_one_or_none()
        if row is None:
            row = CatalogRef(catalog=key[0], ref_code=key[1])
            db.add(row)
            db.flush()
        resolved.append(row)
    return resolved


def _check_grade(db: Session, grade_id: int | None) -> None:
    if grade_id is not None and db.get(Grade, grade_id) is None:
        raise HTTPException(status_code=422, detail=f"Unknown grade_id {grade_id}")


def _check_set(db: Session, set_id: int | None) -> None:
    if set_id is not None and db.get(ItemSet, set_id) is None:
        raise HTTPException(status_code=422, detail=f"Unknown set_id {set_id}")


def _jsonable(value):
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, uuid.UUID):
        return str(value)
    return value


def record_event(db: Session, item_id: uuid.UUID, action: str, changes: dict | None = None):
    db.add(ItemEvent(item_id=item_id, action=action, changes=changes))


def _latest_value_subquery(column=PriceEstimate.estimated_value):
    return (
        select(column)
        .where(PriceEstimate.item_id == Item.id)
        .order_by(PriceEstimate.fetched_at.desc(), PriceEstimate.id)
        .limit(1)
        .correlate(Item)
        .scalar_subquery()
    )


def _filtered(
    stmt,
    *,
    type: str | None,
    status: str | None,
    strike: str | None,
    country: str | None,
    year: int | None,
    year_min: int | None,
    year_max: int | None,
    nd: bool | None,
    tag: str | None,
    set_id: int | None,
    grade_min: int | None,
    grade_max: int | None,
    value_min: float | None,
    value_max: float | None,
    q: str | None,
    fancy: bool | None = None,
    serial_trait: str | None = None,
    target_reached: bool | None = None,
):
    if type:
        stmt = stmt.where(Item.type == type)
    if status:
        stmt = stmt.where(Item.status == status)
    if strike:
        stmt = stmt.where(Item.strike == strike)
    if country:
        stmt = stmt.where(Item.country.ilike(country))
    if year is not None:
        stmt = stmt.where(Item.year == year)
    if year_min is not None:
        stmt = stmt.where(Item.year >= year_min)
    if year_max is not None:
        stmt = stmt.where(Item.year <= year_max)
    # A null year never matches the year filters above (SQL comparison with
    # NULL), which is what an undated piece should do.
    if nd is not None:
        stmt = stmt.where(Item.year_nd.is_(bool(nd)))
    if tag:
        stmt = stmt.where(Item.tags.any(Tag.name.ilike(tag)))
    if set_id is not None:
        stmt = stmt.where(Item.set_id == set_id)
    if grade_min is not None:
        stmt = stmt.where(Item.grade.has(Grade.rank >= grade_min))
    if grade_max is not None:
        stmt = stmt.where(Item.grade.has(Grade.rank <= grade_max))
    if value_min is not None:
        stmt = stmt.where(_latest_value_subquery() >= value_min)
    if value_max is not None:
        stmt = stmt.where(_latest_value_subquery() <= value_max)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(
            or_(
                Item.notes.ilike(like),
                Item.series.ilike(like),
                Item.variety.ilike(like),
                Item.country.ilike(like),
                Item.denomination.ilike(like),
                Item.cert_number.ilike(like),
                Item.serial_number.ilike(like),
                Item.prefix_block.ilike(like),
                Item.issuer.ilike(like),
                Item.charter_number.ilike(like),
                Item.bank_city.ilike(like),
                Item.printer.ilike(like),
                Item.catalog_refs.any(CatalogRef.ref_code.ilike(like)),
                Item.tags.any(Tag.name.ilike(like)),
            )
        )
    if fancy:
        stmt = stmt.where(Item.serial_traits.is_not(None))
    if serial_trait:
        stmt = stmt.where(Item.serial_traits.like(f"%,{serial_trait},%"))
    if target_reached:
        # As Item.target_gap: the newest estimate, in the item's own currency.
        stmt = stmt.where(
            Item.target_price.is_not(None),
            _latest_value_subquery(PriceEstimate.currency) == Item.currency,
            _latest_value_subquery() <= Item.target_price,
        )
    return stmt


def _list_entry(
    item: Item,
    strategy: str,
    preferred_source: str | None,
    converter: Converter | None,
) -> ItemListEntry:
    primary = next((p for p in item.photos if p.is_primary), None)
    if primary is None and item.photos:
        primary = item.photos[0]
    resolved = pricing.resolve_display_value(item.estimates, strategy, preferred_source, converter)
    return ItemListEntry(
        **ItemOut.model_validate(item).model_dump(),
        primary_photo_key=primary.file_key if primary else None,
        primary_thumb_key=(primary.thumb_key or primary.file_key) if primary else None,
        latest_value=float(resolved[0]) if resolved else None,
        latest_value_currency=resolved[1] if resolved else None,
        latest_value_source=resolved[2] if resolved else None,
    )


def _value_settings(db: Session) -> tuple[str, str | None, Converter | None]:
    """The value-resolution strategy, plus a Converter only when the
    strategy needs one: "latest"/"preferred_source" never build one, so the
    default path costs no extra currency lookups."""
    strategy = str(app_settings.get_setting(db, "value_strategy"))
    preferred_source = app_settings.get_setting(db, "preferred_source")
    converter = Converter(db, app_settings.display_currency(db)) if strategy == "average" else None
    return strategy, preferred_source, converter


def filter_query(
    type: str | None = Query(default=None, pattern="^(coin|note)$"),
    status: str | None = Query(default=None, pattern="^(owned|sold|wishlist)$"),
    strike: str | None = Query(default=None, pattern="^(business|proof|specimen)$"),
    country: str | None = None,
    year: int | None = None,
    year_min: int | None = None,
    year_max: int | None = None,
    nd: bool | None = None,
    tag: str | None = None,
    set_id: int | None = None,
    grade_min: int | None = Query(default=None, ge=1, le=70),
    grade_max: int | None = Query(default=None, ge=1, le=70),
    value_min: float | None = Query(default=None, ge=0),
    value_max: float | None = Query(default=None, ge=0),
    q: str | None = None,
    fancy: bool | None = None,
    serial_trait: str | None = Query(default=None, max_length=20),
    target_reached: bool | None = None,
) -> dict:
    """The list filters, shared by list and both exports via Depends."""
    if serial_trait is not None and serial_trait not in serials.TRAITS:
        raise HTTPException(
            status_code=422,
            detail=f"Unknown serial trait {serial_trait!r}; expected one of "
            + ", ".join(serials.TRAITS),
        )
    return {
        "type": type,
        "status": status,
        "strike": strike,
        "country": country,
        "year": year,
        "year_min": year_min,
        "year_max": year_max,
        "nd": nd,
        "tag": tag,
        "set_id": set_id,
        "grade_min": grade_min,
        "grade_max": grade_max,
        "value_min": value_min,
        "value_max": value_max,
        "q": q,
        "fancy": fancy,
        "serial_trait": serial_trait,
        "target_reached": target_reached,
    }


@router.get("", response_model=ItemList)
@permission("read")
def list_items(
    filters: dict = Depends(filter_query),
    sort: str = "-created_at",
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
):
    field = sort.lstrip("-")
    descending = sort.startswith("-")
    if field == "grade":
        order = Grade.rank.desc() if descending else Grade.rank.asc()
    elif field == "value":
        # The newest estimate, the same subquery the value filters use.
        value_column = _latest_value_subquery()
        order = value_column.desc() if descending else value_column.asc()
    elif field in SORTABLE:
        order = SORTABLE[field].desc() if descending else SORTABLE[field].asc()
    else:
        raise HTTPException(status_code=422, detail=f"Unknown sort field: {field}")
    # `IS NULL` sorts false first on SQLite and Postgres alike, so empty rows
    # come last whichever way the column runs.
    if field == "value":
        ordering = (_latest_value_subquery().is_(None), order)
    elif field in NULLS_LAST:
        ordering = (SORTABLE[field].is_(None), order)
    else:
        ordering = (order,)

    stmt = _filtered(select(Item), **filters)
    if field == "grade":
        stmt = stmt.outerjoin(Grade, Item.grade_id == Grade.id)
    total = db.execute(select(func.count()).select_from(stmt.subquery())).scalar_one()
    rows = (
        db.execute(
            stmt.options(*ITEM_LOAD).order_by(*ordering, Item.id).limit(limit).offset(offset)
        )
        .scalars()
        .all()
    )
    strategy, preferred_source, converter = _value_settings(db)
    items = [_list_entry(i, strategy, preferred_source, converter) for i in rows]
    return ItemList(items=items, total=total, limit=limit, offset=offset)


@router.get("/export.csv")
@permission("admin", fresh=True)
def export_csv(
    request: Request, filters: dict = Depends(filter_query), db: Session = Depends(get_db)
):
    # Audited first: the commit would otherwise expire the rows being streamed.
    events.record(
        db, request, "export_downloaded", target="csv", alert="The collection was exported (CSV)."
    )
    rows = (
        db.execute(_filtered(select(Item), **filters).options(*ITEM_LOAD).order_by(Item.created_at))
        .scalars()
        .all()
    )
    strategy, preferred_source, converter = _value_settings(db)

    def generate():
        buf = io.StringIO()
        writer = csv.writer(buf)
        # A byte-order mark tells Excel the file is UTF-8 (otherwise "Schön"
        # opens as "SchÃ¶n"); Cabinet's import reads it either way.
        buf.write("﻿")
        writer.writerow(CSV_COLUMNS)
        yield buf.getvalue()  # the header, even when nothing matches
        buf.seek(0)
        buf.truncate(0)
        for item in rows:
            writer.writerow(_export_row(item, strategy, preferred_source, converter))
            yield buf.getvalue()
            buf.seek(0)
            buf.truncate(0)

    return StreamingResponse(
        generate(),
        media_type="text/csv",
        headers={
            "Content-Disposition": 'attachment; filename="cabinet-items.csv"',
            "Cache-Control": "private, no-store",
        },
    )


@router.get("/export.xlsx")
@permission("admin", fresh=True)
def export_xlsx(
    request: Request, filters: dict = Depends(filter_query), db: Session = Depends(get_db)
):
    # Audited first: the commit would otherwise expire the rows being streamed.
    events.record(
        db,
        request,
        "export_downloaded",
        target="xlsx",
        alert="The collection was exported (Excel).",
    )
    rows = (
        db.execute(_filtered(select(Item), **filters).options(*ITEM_LOAD).order_by(Item.created_at))
        .scalars()
        .all()
    )

    strategy, preferred_source, converter = _value_settings(db)
    wb = Workbook()
    ws = wb.active
    ws.title = "Collection"
    ws.append(CSV_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for item in rows:
        row = _export_row(item, strategy, preferred_source, converter)
        ws.append([str(v) if isinstance(v, uuid.UUID) else v for v in row])
    for column, header in zip(ws.columns, CSV_COLUMNS, strict=True):
        width = max(len(header), *(len(str(c.value or "")) for c in column))
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="cabinet-items.xlsx"',
            "Cache-Control": "private, no-store",
        },
    )


def _export_row(
    item: Item,
    strategy: str,
    preferred_source: str | None,
    converter: Converter | None,
) -> list:
    resolved = pricing.resolve_display_value(item.estimates, strategy, preferred_source, converter)
    flag = lambda value: "true" if value else ""  # noqa: E731
    return [
        item.id,
        item.type,
        item.status,
        item.country,
        item.denomination,
        "" if item.year is None else item.year,
        flag(item.year_nd),
        item.mint_mark or "",
        item.series or "",
        item.variety or "",
        item.strike,
        item.composition or "",
        item.weight_g or "",
        item.fineness or "",
        item.diameter_mm or "",
        item.thickness_mm or "",
        item.width_mm or "",
        item.height_mm or "",
        item.edge or "",
        item.shape or "",
        "" if item.mintage is None else item.mintage,
        "" if item.die_axis is None else item.die_axis,
        item.struck_calendar or "",
        item.struck_year or "",
        item.struck_era or "",
        item.grade.scale if item.grade else "",
        item.grade.code if item.grade else "",
        flag(item.grade_plus),
        flag(item.grade_star),
        "|".join(item.designations or []),
        item.grade_details or "",
        item.cac_sticker or "",
        item.cert_service or "",
        item.cert_number or "",
        "" if item.pcgs_population is None else item.pcgs_population,
        "" if item.pcgs_pop_higher is None else item.pcgs_pop_higher,
        item.serial_number or "",
        item.prefix_block or "",
        item.signatures or "",
        item.issuer or "",
        flag(item.replacement_note),
        item.charter_number or "",
        item.bank_city or "",
        item.bank_state or "",
        item.plate_position or "",
        item.printer or "",
        item.watermark or "",
        item.demonetized_on or "",
        item.quantity,
        item.acquisition_date or "",
        item.acquisition_price or "",
        item.acquisition_fees or "",
        item.spot_at_purchase or "",
        item.currency,
        item.acquired_from or "",
        item.storage_location or "",
        item.sold_date or "",
        item.sold_price or "",
        item.sold_fees or "",
        item.sold_to or "",
        item.target_price or "",
        item.priority or "",
        item.notes or "",
        "|".join(t.name for t in item.tags),
        "|".join(f"{r.catalog}:{r.ref_code}" for r in item.catalog_refs),
        item.set.name if item.set else "",
        json.dumps(item.custom_fields) if item.custom_fields else "",
        resolved[0] if resolved else "",
        resolved[1] if resolved else "",
        item.created_at.isoformat(),
    ]


def _resolve_grade(db: Session, scale: str, code: str) -> tuple[Grade, str | None, bool]:
    """A grade code from a CSV, read the way holders write it: a trailing "+"
    is a plus grade, and on the Sheldon scale PR-/PF-/SP- mean the same number
    with a proof or specimen strike. Returns (grade, strike or None, plus)."""
    plus = code.endswith("+")
    code = code.rstrip("+").strip()
    strike = None
    stmt = select(Grade).where(Grade.scale == scale, Grade.code == code)
    prefix, _, number = code.partition("-")
    if scale == "sheldon" and prefix.upper() in STRIKE_PREFIXES and number.isdigit():
        strike = STRIKE_PREFIXES[prefix.upper()]
        stmt = select(Grade).where(Grade.scale == scale, Grade.rank == int(number))
    grade = db.execute(stmt).scalars().first()
    if grade is None:
        raise ValueError(f"Unknown grade {code!r} on scale {scale!r}")
    return grade, strike, plus


def _row_to_payload(row: dict, db: Session) -> tuple[ItemCreate, int | None]:
    """Map a CSV row (export format) to an ItemCreate + resolved grade_id."""
    data: dict = {
        k: v for k, v in row.items() if k and k not in IMPORT_IGNORED and v not in (None, "")
    }
    # "ND", "ND (1951)", or an empty cell (an older export has no `year_nd`).
    year, undated = importing.parse_year(data.pop("year", None))
    if year is not None:
        data["year"] = year
    if undated and not data.get("year_nd"):
        data["year_nd"] = True
    grade_id = None
    scale = data.pop("grade_scale", "").strip().lower() or "sheldon"
    code = data.pop("grade", "").strip()
    if code:
        grade, strike, plus = _resolve_grade(db, scale, code)
        grade_id = grade.id
        if strike:
            data.setdefault("strike", strike)
        if plus:
            data.setdefault("grade_plus", True)
    designations = [d for d in data.pop("designations", "").split("|") if d.strip()]
    if designations:
        data["designations"] = designations
    tags = [t for t in data.pop("tags", "").split("|") if t.strip()]
    refs = []
    for chunk in data.pop("catalog_refs", "").split("|"):
        if not chunk.strip():
            continue
        if ":" not in chunk:
            raise ValueError(f"Bad catalog ref {chunk!r}; expected 'catalog:code'")
        catalog, ref_code = chunk.split(":", 1)
        refs.append({"catalog": catalog.strip(), "ref_code": ref_code.strip()})

    set_id = None
    set_name = data.pop("set", "").strip()
    if set_name:
        existing = db.execute(select(ItemSet).where(ItemSet.name == set_name)).scalar_one_or_none()
        if existing is None:
            existing = ItemSet(name=set_name)
            db.add(existing)
            db.flush()
        set_id = existing.id

    raw_cf = data.pop("custom_fields", "")
    if raw_cf:
        try:
            data["custom_fields"] = json.loads(raw_cf)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Bad custom_fields JSON: {exc}") from exc

    payload = ItemCreate(**data, tags=tags, catalog_refs=refs, set_id=set_id)
    return payload, grade_id


def _build_item(db: Session, payload: ItemCreate, grade_id: int | None = None) -> Item:
    data = payload.model_dump(exclude={"tags", "catalog_refs", "grade_id"})
    item = Item(**data)
    item.serial_traits = serials.stored_traits(item.serial_number, item.replacement_note)
    if item.pcgs_population is not None or item.pcgs_pop_higher is not None:
        item.population_as_of = datetime.now(timezone.utc)
    # A purchase-day spot price that arrives with the item was typed in; the
    # backfill is the only thing that writes "auto".
    if item.spot_at_purchase is not None:
        item.spot_at_purchase_source = "manual"
    item.grade_id = grade_id if grade_id is not None else payload.grade_id
    _check_grade(db, item.grade_id)
    _check_set(db, item.set_id)
    item.tags = resolve_tags(db, payload.tags)
    item.catalog_refs = resolve_catalog_refs(db, payload.catalog_refs)
    return item


@router.post("", response_model=ItemOut, status_code=201)
@permission("write")
def create_item(payload: ItemCreate, db: Session = Depends(get_db)):
    item = _build_item(db, payload)
    db.add(item)
    db.flush()
    record_event(db, item.id, "created")
    db.commit()
    return get_item_or_404(db, item.id, load_related=True)


@router.post("/run", response_model=RunResult, status_code=201)
@permission("write")
def add_run(payload: RunCreate, db: Session = Depends(get_db)):
    """One item per chosen issue of a Numista type: a date/mint run in one
    request. The type fills each item as "Fill from Numista" would; `shared`
    supplies what they have in common; issues already owned are skipped."""
    try:
        found = numista.catalogue_type(db, payload.type_id)
    except pricing.NotApplicable as exc:
        raise HTTPException(422, str(exc)) from None
    except numista.CatalogueNotFound as exc:
        raise HTTPException(404, str(exc)) from None
    except pricing.SourceUnavailable as exc:
        raise HTTPException(502, str(exc)) from None
    fields = {k: v for k, v in found["fields"].items() if k not in ("year", "mintage")}
    shared = payload.shared.model_dump(exclude_none=True)
    shared["currency"] = shared["currency"].upper()
    owned = (
        checklists.owned_by_issue(db, catalog="numista", ref=f"N#{payload.type_id}")
        if payload.skip_owned
        else {}
    )
    created, skipped, seen = [], 0, set()
    for issue in payload.issues:
        key = (issue.year, checklists.norm(issue.mint_mark))
        if key in owned or key in seen:
            skipped += 1
            continue
        seen.add(key)
        try:
            item_payload = ItemCreate(
                **{
                    **fields,
                    **shared,
                    "year": issue.year,
                    "year_nd": issue.nd,
                    "mint_mark": issue.mint_mark or None,
                    "mintage": issue.mintage,
                    "catalog_refs": found["catalog_refs"],
                }
            )
        except ValidationError as exc:
            err = exc.errors()[0]
            where = ".".join(str(p) for p in err.get("loc", ()))
            label = issue.year if issue.year is not None else "ND"
            raise HTTPException(422, f"{label}: {where}: {err['msg']}") from None
        item = _build_item(db, item_payload)
        db.add(item)
        db.flush()
        record_event(db, item.id, "created")
        created.append(item.id)
    db.commit()
    return RunResult(created=len(created), skipped=skipped, item_ids=created)


@router.get("/similar", response_model=list[SimilarItem])
@permission("read")
def similar_items(
    country: str | None = Query(default=None, max_length=100),
    denomination: str | None = Query(default=None, max_length=100),
    year: int | None = Query(default=None, ge=-5000, le=3000),
    nd: bool = False,
    mint_mark: str | None = Query(default=None, max_length=10),
    cert_number: str | None = Query(default=None, max_length=50),
    ref: list[str] = Query(default=[], description="catalog:code, repeatable"),
    exclude: uuid.UUID | None = None,
    db: Session = Depends(get_db),
):
    """Items that look like the one being entered (the same cert number, the
    same catalogue reference, or the same country, denomination, year, and
    mint mark), trash included, so a restore can replace a re-entry."""
    refs = [tuple(r.split(":", 1)) for r in ref if ":" in r]
    return [
        SimilarItem(
            id=item.id,
            label=item.label,
            grade_label=item.grade_label,
            status=item.status,
            in_trash=item.deleted_at is not None,
            reason=reason,
        )
        for item, reason in duplicates.find_similar(
            db,
            country=country,
            denomination=denomination,
            year=year,
            nd=nd,
            mint_mark=mint_mark,
            cert_number=cert_number,
            refs=refs,
            exclude_id=exclude,
        )
    ]


@router.get("/{item_id}", response_model=ItemDetail)
@permission("read")
def get_item(item_id: uuid.UUID, db: Session = Depends(get_db)):
    """One item with its photos, values, documents, and sales, including an
    item in the trash (`deleted_at` set), which is read-only until restored."""
    return get_item_or_404(db, item_id, load_related=True, include_deleted=True)


@router.patch("/{item_id}", response_model=ItemOut)
@permission("write")
def update_item(item_id: uuid.UUID, payload: ItemUpdate, db: Session = Depends(get_db)):
    item = get_item_or_404(db, item_id, load_related=True)
    fields = payload.model_dump(exclude_unset=True)
    changes: dict = {}
    if "tags" in fields:
        old_tags = sorted(t.name for t in item.tags)
        item.tags = resolve_tags(db, fields.pop("tags") or [])
        new_tags = sorted(t.name for t in item.tags)
        if old_tags != new_tags:
            changes["tags"] = [old_tags, new_tags]
    if "catalog_refs" in fields:
        item.catalog_refs = resolve_catalog_refs(db, payload.catalog_refs or [])
        fields.pop("catalog_refs")
    if "grade_id" in fields:
        _check_grade(db, fields["grade_id"])
    if "set_id" in fields:
        _check_set(db, fields["set_id"])
    _apply_struck_date(item, fields)
    for field, value in fields.items():
        old = getattr(item, field)
        if old != value:
            changes[field] = [_jsonable(old), _jsonable(value)]
        setattr(item, field, value)
    _sync_derived(item, changes)
    if changes:
        record_event(db, item.id, "updated", changes)
    db.commit()
    return get_item_or_404(db, item_id, load_related=True)


def _apply_struck_date(item: Item, fields: dict) -> None:
    """The era rule, checked against the item as it will be; `year: null`
    beside a date as struck is answered with the converted year, and on an
    undated piece it stays empty. A year already there is left alone."""
    merged = {
        key: fields.get(key, getattr(item, key))
        for key in ("struck_calendar", "struck_year", "struck_era")
    }
    try:
        calendars.check_era(merged["struck_calendar"], merged["struck_era"])
        if (
            "year" in fields
            and fields["year"] is None
            and merged["struck_calendar"]
            and merged["struck_year"] is not None
        ):
            fields["year"] = calendars.to_gregorian(
                merged["struck_calendar"], merged["struck_year"], merged["struck_era"]
            )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from None
    # The year rule, against the item as it will be.
    if fields.get("year", item.year) is None and not fields.get("year_nd", item.year_nd):
        raise HTTPException(status_code=422, detail=YEAR_REQUIRED)


def _sync_derived(item: Item, changed) -> None:
    """The server-set fields, after an edit: the serial's traits, when the
    population figures last changed, and where a purchase-day spot came from."""
    if "serial_number" in changed or "replacement_note" in changed:
        item.serial_traits = serials.stored_traits(item.serial_number, item.replacement_note)
    if "pcgs_population" in changed or "pcgs_pop_higher" in changed:
        has_any = item.pcgs_population is not None or item.pcgs_pop_higher is not None
        item.population_as_of = datetime.now(timezone.utc) if has_any else None
    if "spot_at_purchase" in changed:
        item.spot_at_purchase_source = "manual" if item.spot_at_purchase is not None else None


@router.get("/{item_id}/history", response_model=list[EventOut])
@permission("read")
def item_history(item_id: uuid.UUID, db: Session = Depends(get_db)):
    get_item_or_404(db, item_id)
    return (
        db.execute(
            select(ItemEvent).where(ItemEvent.item_id == item_id).order_by(ItemEvent.id.desc())
        )
        .scalars()
        .all()
    )


@router.post("/bulk", response_model=BulkResult)
@permission("write")
def bulk_update(payload: BulkUpdate, db: Session = Depends(get_db)):
    """Apply the same changes to many items at once: scalar field updates via
    `set` (tags/refs inside it are ignored), plus add_tags / remove_tags."""
    items = (
        db.execute(select(Item).where(Item.id.in_(payload.ids)).options(selectinload(Item.tags)))
        .scalars()
        .all()
    )
    if len(items) != len(set(payload.ids)):
        raise HTTPException(status_code=404, detail="One or more items not found")

    fields = {}
    if payload.set is not None:
        fields = payload.set.model_dump(exclude_unset=True)
        fields.pop("tags", None)
        fields.pop("catalog_refs", None)
        # ND is per piece, and a bulk year stays as it is (no struck-date check),
        # so a bulk edit can't empty the year either: that needs ND.
        fields.pop("year_nd", None)
        # A purchase-day spot price is per piece, and looked up per piece.
        fields.pop("spot_at_purchase", None)
        # Note details and demonetization are per piece, not bulk-editable.
        for key in ("width_mm", "height_mm", "printer", "watermark", "demonetized_on"):
            fields.pop(key, None)
        if "year" in fields and fields["year"] is None:
            raise HTTPException(status_code=422, detail=YEAR_REQUIRED)
        if "grade_id" in fields:
            _check_grade(db, fields["grade_id"])
        if "set_id" in fields:
            _check_set(db, fields["set_id"])

    add = resolve_tags(db, payload.add_tags)
    remove_names = {n.strip() for n in payload.remove_tags if n.strip()}
    summary = {f: [None, _jsonable(v)] for f, v in fields.items()}
    if payload.add_tags:
        summary["add_tags"] = [None, payload.add_tags]
    if remove_names:
        summary["remove_tags"] = [None, sorted(remove_names)]
    for item in items:
        touched = {f for f, value in fields.items() if getattr(item, f) != value}
        for field, value in fields.items():
            setattr(item, field, value)
        _sync_derived(item, touched)
        if add:
            existing = {t.name for t in item.tags}
            item.tags.extend(t for t in add if t.name not in existing)
        if remove_names:
            item.tags = [t for t in item.tags if t.name not in remove_names]
        if summary:
            record_event(db, item.id, "updated", {"bulk": [None, True], **summary})
    db.commit()
    return BulkResult(updated=len(items))


@router.post("/{item_id}/clone", response_model=ItemOut, status_code=201)
@permission("write")
def clone_item(item_id: uuid.UUID, db: Session = Depends(get_db)):
    """Copy an item's fields (not photos or estimates) to speed up entering
    similar pieces."""
    source = get_item_or_404(db, item_id, load_related=True)
    copy = Item(
        **{
            col.name: getattr(source, col.name)
            for col in Item.__table__.columns
            if col.name
            not in ("id", "created_at", "updated_at", "import_source", "import_key", "deleted_at")
        }
    )
    # The copy's population figures are set now, so they are dated now.
    if copy.population_as_of is not None:
        copy.population_as_of = datetime.now(timezone.utc)
    copy.tags = list(source.tags)
    copy.catalog_refs = list(source.catalog_refs)
    db.add(copy)
    db.flush()
    record_event(db, copy.id, "created", {"cloned_from": [None, str(source.id)]})
    db.commit()
    return get_item_or_404(db, copy.id, load_related=True)


@router.delete("/{item_id}", status_code=204)
@permission("write")
def delete_item(
    item_id: uuid.UUID, request: Request, permanent: bool = False, db: Session = Depends(get_db)
):
    """Move the item to the trash, from where it can be restored. With
    `?permanent=true`, or for an item already in the trash, delete it for
    good, with its photos, values, history, and documents no other item holds:
    that needs the admin and a recent password."""
    item = get_item_or_404(db, item_id, include_deleted=True)
    if permanent or item.deleted_at is not None:
        require(request, "admin", fresh=True)
        trash.purge(db, item)
    else:
        trash.move_to_trash(db, [item])


@router.post("/{item_id}/restore", response_model=ItemOut)
@permission("write")
def restore_item(item_id: uuid.UUID, db: Session = Depends(get_db)):
    """Take an item out of the trash, as it was."""
    item = get_item_or_404(db, item_id, include_deleted=True)
    trash.restore(db, [item])
    return get_item_or_404(db, item_id, load_related=True)
